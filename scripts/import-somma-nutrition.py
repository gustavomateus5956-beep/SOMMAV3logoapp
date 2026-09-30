"""Offline, deterministic TACO importer. Requires Python 3.10+ and openpyxl 3.1.5.

Run from any directory; --check verifies committed outputs without writing.
Only the pinned, attributed official TACO workbook is eligible for distribution.
POF is deliberately not accepted by this importer.
"""
import argparse
import collections
import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'src/data/nutrition'
SOURCE = ROOT / 'data-sources/nutrition/taco-4-2011.xlsx'
MANIFEST = ROOT / 'data-sources/nutrition/sources.json'
VERSION = '1.0.0'
CATEGORIES = ['Cereais e derivados', 'Verduras, hortaliças e derivados', 'Frutas e derivados',
              'Gorduras e óleos', 'Pescados e frutos do mar', 'Carnes e derivados',
              'Leite e derivados', 'Bebidas (alcoólicas e não alcoólicas)', 'Ovos e derivados',
              'Produtos açucarados', 'Miscelâneas', 'Outros alimentos industrializados',
              'Alimentos preparados', 'Leguminosas e derivados', 'Nozes e sementes']
MACROS = {3: 'calories', 5: 'protein', 6: 'fat', 7: 'cholesterol', 8: 'carbohydrates',
          9: 'fiber', 11: 'calcium', 16: 'iron', 17: 'sodium', 18: 'potassium'}
MICROS = {12: ('magnesium','mg'), 14: ('manganese','mg'), 15: ('phosphorus','mg'),
          19: ('copper','mg'), 20: ('zinc','mg'), 21: ('retinol','µg'),
          22: ('retinolEquivalent','µg'), 23: ('retinolActivityEquivalent','µg'),
          24: ('thiamin','mg'), 25: ('riboflavin','mg'), 26: ('pyridoxine','mg'),
          27: ('niacin','mg'), 28: ('vitaminC','mg')}
# Editorial wording only; original names remain searchable and attached to each record.
NAMES = {
 'Arroz, integral, cozido':'Arroz integral cozido',
 'Arroz, integral, cru':'Arroz integral cru',
 'Arroz, tipo 1, cozido':'Arroz branco tipo 1 cozido',
 'Arroz, tipo 1, cru':'Arroz branco tipo 1 cru',
 'Arroz, tipo 2, cozido':'Arroz branco tipo 2 cozido',
 'Arroz, tipo 2, cru':'Arroz branco tipo 2 cru',
 'Feijão, carioca, cozido':'Feijão carioca cozido',
 'Feijão, preto, cozido':'Feijão preto cozido',
 'Pão, francês':'Pão francês',
 'Pão, trigo, francês':'Pão francês de trigo',
 'Pão, aveia, forma':'Pão de forma de aveia',
 'Pão, glúten, forma':'Pão de forma com glúten',
 'Pão, milho, forma':'Pão de forma de milho',
 'Pão, trigo, forma, integral':'Pão de forma integral de trigo',
 'Pão, trigo, sovado':'Pão sovado de trigo',
 'Pão, de queijo, assado':'Pão de queijo assado',
 'Farinha, de mandioca, crua':'Farinha de mandioca crua',
 'Mandioca, cozida':'Mandioca cozida',
 'Mandioca, farofa, temperada':'Farofa de mandioca temperada',
 'Batata, doce, crua':'Batata-doce crua',
 'Batata, doce, cozida':'Batata-doce cozida',
 'Batata, inglesa, cozida':'Batata inglesa cozida',
 'Banana, prata, crua':'Banana-prata crua',
 'Banana, nanica, crua':'Banana-nanica crua',
 'Mamão, Formosa, cru':'Mamão formosa cru',
 'Ovo, de galinha, inteiro, cozido/10minutos':'Ovo de galinha inteiro cozido por 10 minutos',
 'Ovo, de galinha, clara, cozida/10minutos':'Clara de ovo de galinha cozida por 10 minutos',
 'Ovo, de galinha, gema, cozida/10minutos':'Gema de ovo de galinha cozida por 10 minutos',
 'Frango, peito, sem pele, grelhado':'Peito de frango sem pele grelhado',
 'Carne, bovina, patinho, sem gordura, grelhado':'Patinho bovino sem gordura grelhado',
 'Carne, bovina, contra-filé, sem gordura, grelhado':'Contrafilé bovino sem gordura grelhado',
 'Leite, de vaca, integral':'Leite de vaca integral',
 'Leite, de vaca, integral, pó':'Leite de vaca integral em pó',
 'Leite, de vaca, desnatado, pó':'Leite de vaca desnatado em pó',
 'Queijo, minas, frescal':'Queijo minas frescal',
 'Cuscuz, de milho, cozido com sal':'Cuscuz de milho cozido com sal',
 'Café, infusão 10%':'Café em infusão a 10%',
}

def normalize(text):
    return re.sub(r'\s+', ' ', ''.join(c for c in unicodedata.normalize('NFD', text)
                                     if unicodedata.category(c) != 'Mn').lower()).strip()

def display_name(original):
    if original in NAMES:
        return NAMES[original]
    text = re.sub(r',\s*', ' ', original)
    text = text.replace('mingnon', 'mignon').replace('contra-filé', 'contrafilé')
    text = re.sub(r'\s+', ' ', text).strip()
    return text[:1].upper() + text[1:]

def aliases_for(original, name):
    aliases = [original]
    # Substitutions preserve the full preparation, variety and qualifiers.
    replacements = [('Mandioca','Aipim'), ('Mandioca','Macaxeira'),
                    ('mandioca','aipim'), ('mandioca','macaxeira'),
                    ('Abóbora','Jerimum'), ('Couve-flor','Couve flor'),
                    ('Pão francês','Pão de sal'), ('Tangerina','Mexerica')]
    for old,new in replacements:
        if old in name:
            aliases.append(name.replace(old,new))
    result = []
    for alias in aliases:
        if alias != name and alias not in result:
            result.append(alias)
    return result

def parse_value(value):
    if isinstance(value, (int,float)) and not isinstance(value,bool):
        if not math.isfinite(value) or value < 0:
            return None, 'invalidSourceValue'
        return value, None
    marker = str(value).strip() if value is not None else ''
    statuses = {'': 'notAnalyzed', 'Tr':'trace', 'NA':'notApplicable', '*':'underReview'}
    if marker in statuses:
        return None, statuses[marker]
    # Never guess a malformed cell such as ",0,02".
    return None, 'invalidSourceValue'

def build():
    manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
    source = next(s for s in manifest['sources'] if s['id']=='taco-4-2011')
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if digest != source['sha256'] or source['redistribution'] != 'allowed-with-attribution':
        raise ValueError('Source checksum or redistribution policy mismatch')
    book = openpyxl.load_workbook(SOURCE, data_only=True)
    main, fats = book.worksheets[:2]
    assert main.title == 'CMVCol taco3' and fats.title == 'AGtaco3'
    assert main.cell(3,4).value == '(kcal)' and main.cell(3,6).value == '(g)'
    assert main.cell(3,18).value == '(mg)' and main.cell(3,22).value == '(mcg)'
    fat_rows = {}
    for rownum,row in enumerate(fats.values,1):
        if isinstance(row[0],int):
            if row[0] in fat_rows: raise ValueError('Duplicate fatty-acid ID')
            fat_rows[row[0]] = (rownum,row)
    foods, anomalies = [], []
    category = None
    for rownum,row in enumerate(main.values,1):
        if row[0] in CATEGORIES:
            category = row[0]
        if not isinstance(row[0],int): continue
        number, original = row[:2]
        raw_name = original
        if number == 540:
            assert original == 'L'
            original = 'Feijoada'
            anomalies.append({'id':'somma:taco:4:0540','field':'name','cell':'CMVCol taco3!B620',
                              'rawValue':'L','correctedValue':'Feijoada',
                              'evidence':'Official TACO PDF page 63 (printed 60), record 540; AGtaco3 record 540 agrees',
                              'action':'documented correction from the official PDF; no inferred identity'})
        assert category and row[13] == number
        food_id = f'somma:taco:4:{number:04d}'
        nutrition, statuses = {}, {}
        locations = {'name': f'{main.title}!B{rownum}' if number != 540 else 'taco-4-2011.pdf#page=63 (record 540)'}
        def add(key, value, col, sheet, row_number, unit=None):
            number_value, status = parse_value(value)
            locations[key] = f'{sheet}!{openpyxl.utils.get_column_letter(col+1)}{row_number}'
            if status:
                statuses[key] = status
                if status == 'invalidSourceValue':
                    anomalies.append({'id':food_id,'nutrient':key,'cell':locations[key],
                                      'rawValue':value,'action':'omitted; no guessed correction'})
            elif unit:
                nutrition.setdefault('micronutrients',{})[key] = {'value':number_value,'unit':unit}
            else: nutrition[key] = number_value
        for col,key in MACROS.items(): add(key,row[col],col,main.title,rownum)
        for col,(key,unit) in MICROS.items(): add(key,row[col],col,main.title,rownum,unit)
        if number in fat_rows:
            fatrownum,fatrow = fat_rows[number]
            assert normalize(fatrow[1]).replace(' ','') == normalize(original).replace(' ','') and fatrow[12] == number
            add('saturatedFat',fatrow[2],2,fats.title,fatrownum)
        name = display_name(original)
        foods.append({'id':food_id,'name':name,'aliases':aliases_for(original,name),
            'category':normalize(category).replace(' ','-').replace(',','').replace('(','').replace(')',''),
            'nutritionPer100g':nutrition,
            'portions':[{'id':food_id+':100g','label':'100 g de parte comestível (referência)',
                         'amount':100,'unit':'g','grams':100}],
            'householdMeasures':[], 'source':'taco',
            'metadata':{'sourceRecordId':str(number),'sourceVersion':'4ª edição revisada e ampliada, 2011',
                        'sourceUrl':source['url'],'retrievedAt':source['retrievedAt'],
                        'originalName':raw_name,'sourceLicenseId':source['id'],
                        'sourceSha256':digest,'sourceLocations':locations,
                        'nutrientStatus':statuses,
                        'notes':'Fonte: NEPA/UNICAMP, TACO, 4. ed., 2011. Valores por 100 g de parte comestível. A porção de referência não é recomendação de consumo.'}})
    assert len(foods)==597 and len({f['id'] for f in foods})==597
    assert {int(f['metadata']['sourceRecordId']) for f in foods}==set(range(1,598))
    assert set(fat_rows).issubset({int(f['metadata']['sourceRecordId']) for f in foods})
    groups=collections.Counter(f['category'] for f in foods)
    taxonomy={'version':VERSION,'categories':[{'id':normalize(c).replace(' ','-').replace(',','').replace('(','').replace(')',''),
                                              'label':c} for c in CATEGORIES]}
    index={'version':VERSION,'normalization':'NFD; remove combining accents; lowercase; collapse whitespace',
           'entries':[{'id':f['id'],'category':f['category'],
                       'text':normalize(' '.join([f['name'],*f['aliases']]))} for f in foods]}
    coverage=collections.Counter()
    missing=collections.Counter()
    names=collections.defaultdict(list)
    for food in foods:
        names[normalize(food['name'])].append(food['id'])
        for k in food['nutritionPer100g']:
            if k!='micronutrients': coverage[k]+=1
        coverage.update(food['nutritionPer100g'].get('micronutrients',{}).keys())
        missing.update(food['metadata']['nutrientStatus'].values())
    summary={'catalogVersion':VERSION,'builtAt':'2026-09-28','basis':'100 g edible portion',
             'sources':{'taco':{'processed':597,'included':597},
                        'pof':{'included':0,'status':'excluded-pending-commercial-redistribution-clearance',
                               'auditedCompositionRows':next(s['auditedRows'] for s in manifest['sources'] if s['id']=='pof-2008-2009-composition'),
                               'auditedHouseholdMeasureRows':next(s['auditedRows'] for s in manifest['sources'] if s['id']=='pof-2008-2009-householdMeasures')}},
             'total':len(foods),'categoryCounts':dict(groups),'nutrientCoverage':dict(coverage),
             'withReference100gPortion':597,'withSourceServingPortion':0,'withHouseholdMeasures':0,
             'unknownStatuses':dict(missing),'duplicatesRemoved':0,
             'sameNormalizedNameGroups':[v for v in names.values() if len(v)>1],
             'sourceAnomalies':anomalies,
             'notImported':['POF composition and household measures: unresolved downstream rights',
                            'TACO amino acids and individual fatty acids: outside current model scope',
                            'transFat and sugars: no invented totals or borrowed values']}
    return {'somma-foods-br.json':foods,'somma-foods-taxonomy.json':taxonomy,
            'somma-foods-search-index.json':index,'somma-foods-summary.json':summary}

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--check',action='store_true')
    args=parser.parse_args()
    outputs=build()
    OUT.mkdir(parents=True,exist_ok=True)
    for filename,data in outputs.items():
        text=json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False)+'\n'
        path=OUT/filename
        if args.check:
            if not path.exists() or path.read_text(encoding='utf-8')!=text:
                raise ValueError(f'Generated artifact is stale: {filename}')
        else: path.write_text(text,encoding='utf-8',newline='\n')
    print(json.dumps(outputs['somma-foods-summary.json'],ensure_ascii=False,indent=2))

if __name__=='__main__': main()
